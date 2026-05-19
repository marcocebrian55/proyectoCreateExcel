import os
import time
import pandas as pd
from flask import Flask, request, send_file, render_template
from ia_processor import extraer_datos_imagen
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment

# Configuración de rutas
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
UPLOAD_FOLDER = os.path.join(BASE_DIR, 'uploads')

if not os.path.exists(UPLOAD_FOLDER):
    os.makedirs(UPLOAD_FOLDER)

app = Flask(__name__)

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload_file():
    archivos = request.files.getlist('file')
    
    if not archivos or archivos[0].filename == '':
        return "Ningún archivo seleccionado", 400
        
    datos_nuevos_lista = []

    for index, file in enumerate(archivos):
        print(f"\n📸 Procesando imagen {index + 1} de {len(archivos)}: {file.filename}")
        
        imagen_path = os.path.join(UPLOAD_FOLDER, file.filename)
        file.save(imagen_path)
        
        datos = extraer_datos_imagen(imagen_path)
        
        if datos:
            datos_nuevos_lista.append(datos)
            print(f"✅ Datos extraídos de {file.filename}")
        else:
            print(f"⚠️ No se pudieron extraer datos de {file.filename}")
            
        if index < len(archivos) - 1:
            print("⏳ Esperando 15 segundos antes de la siguiente foto...")
            time.sleep(15)
    
    if not datos_nuevos_lista:
        return "Error: No se ha podido procesar ninguna imagen.", 500
        
    try:
        df_nuevos = pd.DataFrame(datos_nuevos_lista)
        excel_path = os.path.join(UPLOAD_FOLDER, 'reporte_wizard.xlsx')
        
        # DEFINIMOS EL ORDEN EXACTO DE LAS COLUMNAS (TOT al final)
        letras = ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P']
        columnas_orden = ['FECHA', 'CATEGORIA'] + letras + ['TOT']

        if os.path.exists(excel_path):
            # Leemos saltando las 2 filas de nuestras cabeceras personalizadas
            df_existente = pd.read_excel(excel_path, skiprows=2, names=columnas_orden)
            df_final = pd.concat([df_existente, df_nuevos], ignore_index=True)
        else:
            df_final = df_nuevos
            
        # Orden y Limpieza
        df_final['FECHA'] = pd.to_datetime(df_final['FECHA'], format='%d/%m/%Y', errors='coerce')
        df_final = df_final.drop_duplicates(subset=['FECHA', 'CATEGORIA'], keep='last')
        df_final = df_final.sort_values(by=['FECHA', 'CATEGORIA'], ascending=[True, True])
        df_final['FECHA'] = df_final['FECHA'].dt.strftime('%d/%m/%Y')

        # Asegurar todas las columnas
        for col in columnas_orden:
            if col not in df_final.columns:
                df_final[col] = 0
                
        df_final = df_final[columnas_orden]

        # Guardamos en excel SIN CABECERAS de pandas para pintar nosotros la doble fila
        df_final.to_excel(excel_path, index=False, header=False, startrow=2)
        
        # 5. DISEÑO Y COLORES (Openpyxl) - CLONANDO TU EXCEL ORIGINAL
        wb = load_workbook(excel_path)
        ws = wb.active
        
        cabecera_fila_1 = ['FECHA', 'CATEGORIA', 'A', 'B', 'C', 'D', 'E (Autom)', 'F (Kodiaq)', 'G (Civic+A3)', 'H (Berl)', 'I(A4)', 'J (T-ROC)', 'K (cabrio)', 'L(A6)', 'M (jumpy)', 'N (tucson)', 'O (Ber 7pax)', 'P(caddyM)', 'DISPONIBLE WZ']
        cabecera_fila_2 = ['', '', 'SA', 'SC', 'SG', 'SJ', 'SK', 'MD', 'MC', 'SM', 'MJ', 'EG', 'LL', 'LP', 'LO', 'MN', 'LJ', 'EE', '']

        color_rojo = PatternFill(start_color='E6B8B7', end_color='E6B8B7', fill_type='solid')
        color_naranja = PatternFill(start_color='FCD5B4', end_color='FCD5B4', fill_type='solid')
        color_verde = PatternFill(start_color='D8E4BC', end_color='D8E4BC', fill_type='solid')
        color_gris = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')

        # Pintar Fila 1
        for col_num, valor in enumerate(cabecera_fila_1, 1):
            cell = ws.cell(row=1, column=col_num, value=valor)
            cell.fill = color_gris
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # Pintar Fila 2
        for col_num, valor in enumerate(cabecera_fila_2, 1):
            cell = ws.cell(row=2, column=col_num, value=valor)
            cell.fill = color_gris
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # Colorear datos (empezando en la fila 3)
        for row in ws.iter_rows(min_row=3, min_col=1, max_col=ws.max_column):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    if cell.value <= 0:
                        cell.fill = color_rojo
                    elif 0 < cell.value <= 15:
                        cell.fill = color_naranja
                    else:
                        cell.fill = color_verde
                cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Ancho de columnas adaptado a los nombres largos
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 14

        wb.save(excel_path)
        print(f"🚀 TODO LISTO. Excel actualizado con formato doble cabecera.")
        
        return send_file(excel_path, as_attachment=True)
        
    except Exception as e:
        print(f"❌ Error final: {e}")
        return f"Error al generar el Excel: {e}", 500

if __name__ == '__main__':
    app.run(debug=True, port=5000)